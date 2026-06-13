import os
import re
import uuid
import time
import sqlite3
import threading
import warnings
from datetime import datetime

import numpy as np
import requests
import nltk
from sklearn.metrics.pairwise import cosine_similarity
from flask import Flask, request, jsonify, render_template, send_file

from hazm import Normalizer, POSTagger, word_tokenize, sent_tokenize

# ─────────────────────────── config ───────────────────────────
POS_MODEL_PATH = os.environ.get("POS_MODEL", "/opt/seo-kwr/models/pos_tagger.model")
RERANKER_DB = os.environ.get("RERANKER_DB", "/opt/seo-reranker/data.db")
UPLOAD_DIR = os.environ.get("KWR_UPLOAD_DIR", "/opt/seo-kwr/uploads")

KEYWORD_NUM = 30
BETA = 0.9
MAX_PHRASE_LEN = 5
DEFAULT_MAX_CHARS = 20000          # POS tagging cap (1-core friendly)
MAX_CANDIDATES = 400               # cap candidates/URL (embedding cost)
EMBED_BATCH = 96

JINA_URL = "https://api.jina.ai/v1/embeddings"
JINA_MODEL = "jina-embeddings-v5-text-small"
VOYAGE_URL = "https://api.voyageai.com/v1/embeddings"
VOYAGE_MODEL = "voyage-3.5"

GRAMMERS = [
    "NP: {<NOUN,EZ>?<NOUN.*>}",
    "NP: {<NOUN.*><ADJ.*>?}",
    "NP: {<NOUN,EZ><NOUN,EZ><NOUN.*>}",
    "NP: {<NOUN,EZ><NOUN.*><ADJ.*>}",
]

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 64 * 1024 * 1024  # 64MB uploads
from serpiwi_auth import init_auth
init_auth(app, "Serpiwi · استخراج کلیدواژه")

# load hazm once (shared, single worker)
normalizer = Normalizer()
_tagger = None
_tagger_lock = threading.Lock()


def get_tagger():
    global _tagger
    if _tagger is None:
        with _tagger_lock:
            if _tagger is None:
                _tagger = POSTagger(model=POS_MODEL_PATH)
    return _tagger


# ─────────────────────────── job registry ───────────────────────────
JOBS = {}
JOBS_LOCK = threading.Lock()


def new_job():
    jid = uuid.uuid4().hex[:12]
    with JOBS_LOCK:
        JOBS[jid] = {"status": "running", "progress": 0, "log": [], "results": None,
                     "error": None, "keyword": None}
    return jid


def job_update(jid, **kw):
    with JOBS_LOCK:
        j = JOBS.get(jid)
        if not j:
            return
        line = kw.pop("log", None)
        if line is not None:
            j["log"].append(line)
        j.update(kw)


def job_get(jid):
    with JOBS_LOCK:
        j = JOBS.get(jid)
        return dict(j) if j else None


# ─────────────────────────── NLP core (EmbedRank) ───────────────────────────
_MD_LINK = re.compile(r"!?\[([^\]]*)\]\([^)]*\)")
_MD_SYM = re.compile(r"[*_`#>~|]+")
_URL = re.compile(r"https?://\S+")


def clean_content(text):
    """Strip markdown emphasis/links/urls so they don't pollute keyphrases."""
    text = _MD_LINK.sub(r"\1", text)   # [label](url) / ![alt](img) -> label/alt
    text = _URL.sub(" ", text)
    text = _MD_SYM.sub(" ", text)      # ** __ ` # > ~ | table/heading/bold marks
    return re.sub(r"[ \t]+", " ", text)


def pos_tag_text(text):
    norm = normalizer.normalize(text)
    tokens = [word_tokenize(s) for s in sent_tokenize(norm)]
    return get_tagger().tag_sents(tokens)


def extract_all_candidates(tagged_sents):
    cands = set()
    for grammar in GRAMMERS:
        parser = nltk.RegexpParser(grammar)
        for tree in parser.parse_sents(tagged_sents):
            for sub in tree.subtrees(filter=lambda t: t.label() == "NP"):
                phrase = " ".join(w for w, _ in sub.leaves()).strip()
                if phrase and len(phrase.split()) <= MAX_PHRASE_LEN:
                    cands.add(phrase)
    return np.array(sorted(cands))


def embed_texts(provider, texts, api_key):
    """Embed a list of texts via Jina or Voyage; returns np.array (len(texts), D)."""
    texts = [(t or " ")[:8000] for t in texts]
    out = []
    for i in range(0, len(texts), EMBED_BATCH):
        batch = texts[i:i + EMBED_BATCH]
        backoff = [4, 8, 16, 30]
        attempt = 0
        while True:
            if provider == "voyage":
                resp = requests.post(VOYAGE_URL,
                    headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
                    json={"input": batch, "model": VOYAGE_MODEL}, timeout=90)
            else:
                resp = requests.post(JINA_URL,
                    headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
                    json={"model": JINA_MODEL, "task": "text-matching", "normalized": True, "input": batch},
                    timeout=90)
            if resp.status_code == 429 and attempt < len(backoff):
                time.sleep(backoff[attempt]); attempt += 1; continue
            if not resp.ok:
                label = "Voyage" if provider == "voyage" else "Jina"
                raise RuntimeError(f"{label} embeddings {resp.status_code}: {resp.text[:200]}")
            break
        data = sorted(resp.json()["data"], key=lambda x: x["index"])
        out.extend(d["embedding"] for d in data)
    return np.array(out, dtype=float)


def compute_similarities(vecs, text_vec):
    sim_text = cosine_similarity(vecs, text_vec.reshape(1, -1))
    sim_pair = cosine_similarity(vecs)

    sim_text_n = sim_text / np.max(sim_text)
    std = np.std(sim_text_n) or 1.0
    sim_text_n = 0.5 + (sim_text_n - np.average(sim_text_n)) / std

    np.fill_diagonal(sim_pair, np.nan)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        sim_pair_n = sim_pair / np.nanmax(sim_pair, axis=0)
        denom = np.nanstd(sim_pair_n, axis=0)
        denom[denom == 0] = 1.0
        sim_pair_n = 0.5 + (sim_pair_n - np.nanmean(sim_pair_n, axis=0)) / denom
    return sim_text_n, sim_pair_n


def embed_rank_select(candidates, sim_text, sim_pair, keyword_num=KEYWORD_NUM, beta=BETA):
    n = len(candidates)
    N = min(n, keyword_num)
    selected, unselected = [], list(range(n))
    best = int(np.argmax(sim_text))
    selected.append(best); unselected.remove(best)
    for _ in range(N - 1):
        if not unselected:
            break
        sel = np.array(selected); unsel = np.array(unselected)
        rel = sim_text[unsel, :]
        red = sim_pair[unsel][:, sel]
        if red.ndim == 1:
            red = red[:, np.newaxis]
        red = np.nan_to_num(red, nan=0.0)
        mmr = beta * rel - (1 - beta) * np.max(red, axis=1).reshape(-1, 1)
        bl = int(np.argmax(mmr))
        bg = unselected[bl]
        selected.append(bg); unselected.remove(bg)
    return candidates[selected].tolist()


def extract_keyphrases(text, provider, api_key, keyword_num=KEYWORD_NUM):
    tagged = pos_tag_text(clean_content(text))
    candidates = extract_all_candidates(tagged)
    if len(candidates) == 0:
        return [], 0
    if len(candidates) > MAX_CANDIDATES:
        candidates = candidates[:MAX_CANDIDATES]
    embeds = embed_texts(provider, list(candidates) + [" ".join(candidates)], api_key)
    vecs, text_vec = embeds[:-1], embeds[-1]
    sim_text, sim_pair = compute_similarities(vecs, text_vec)
    return embed_rank_select(candidates, sim_text, sim_pair, keyword_num), len(candidates)


# ─────────────────────────── DB reader ───────────────────────────
def list_db_queries(db_path):
    if not os.path.exists(db_path):
        return []
    conn = sqlite3.connect(db_path)
    try:
        rows = conn.execute("""
            SELECT q.id, q.query, q.timestamp, COUNT(DISTINCT u.id)
            FROM queries q LEFT JOIN urls u ON u.query_id = q.id
            GROUP BY q.id ORDER BY q.id DESC LIMIT 100""").fetchall()
    except sqlite3.Error:
        return []
    finally:
        conn.close()
    return [{"id": r[0], "query": r[1], "timestamp": r[2], "urls": r[3]} for r in rows]


def load_urls_from_db(db_path, query_id=None):
    conn = sqlite3.connect(db_path)
    cur = conn.cursor()
    cur.execute("PRAGMA table_info(urls)")
    cols = [r[1] for r in cur.fetchall()]
    has_full = "full_content" in cols
    sel = "u.full_content" if has_full else "NULL"
    where = "WHERE u.query_id = ?" if query_id else ""
    params = (query_id,) if query_id else ()
    cur.execute(f"""
        SELECT u.id, u.url, u.title, u.rank, q.query, {sel}
        FROM urls u LEFT JOIN queries q ON u.query_id = q.id
        {where} ORDER BY q.timestamp DESC, u.rank""", params)
    rows = cur.fetchall()
    results, seen = [], set()
    for url_id, url, title, rank, query, full in rows:
        if url in seen:
            continue
        seen.add(url)
        if not full or not str(full).strip():
            ch = cur.execute("SELECT chunk_text FROM chunks WHERE url_id=? ORDER BY chunk_index", (url_id,)).fetchall()
            full = "\n".join(r[0] for r in ch if r[0])
        results.append({"url_id": url_id, "url": url, "title": title or "",
                        "rank": rank, "original_query": query or "", "content": full or ""})
    conn.close()
    return results


# ─────────────────────────── job ───────────────────────────
def run_extract(jid, db_path, query_id, provider, api_key, keyword_num, max_chars):
    try:
        job_update(jid, log="بارگذاری URLها از دیتابیس…")
        items = load_urls_from_db(db_path, query_id)
        if not items:
            raise RuntimeError("هیچ داده‌ای در دیتابیس پیدا نشد.")
        job_update(jid, log=f"{len(items)} URL برای استخراج کلیدواژه.")
        results, total = [], len(items)
        for i, it in enumerate(items, 1):
            job_update(jid, progress=int((i - 1) / total * 100),
                       log=f"[{i}/{total}] {it['title'][:50] or it['url'][:50]}")
            content = (it["content"] or "").strip()
            if not content:
                job_update(jid, log="   ⚠ محتوا خالی، رد شد.")
                it["keywords"] = []; results.append(it); continue
            if len(content) > max_chars:
                content = content[:max_chars]
            try:
                kws, ncand = extract_keyphrases(content, provider, api_key, keyword_num)
                it["keywords"] = kws
                job_update(jid, log=f"   ✓ {ncand} کاندیدا → {len(kws)} کلیدواژه")
            except RuntimeError as e:
                it["keywords"] = []
                job_update(jid, log=f"   ✗ {e}")
            results.append(it)
        # frequency across urls
        freq = {}
        for it in results:
            for kw in it.get("keywords", []):
                freq[kw] = freq.get(kw, 0) + 1
        freq_sorted = sorted(freq.items(), key=lambda x: (-x[1], x[0]))
        job_update(jid, status="done", progress=100,
                   results={"items": results, "freq": freq_sorted, "total_urls": len(results)},
                   log="✓ استخراج کلیدواژه کامل شد.")
    except Exception as e:
        job_update(jid, status="error", error=str(e), log=f"✗ {e}")


# ─────────────────────────── routes ───────────────────────────
@app.route("/")
def index():
    return render_template("index.html",
                           has_jina=bool(os.environ.get("JINA_API_KEY")),
                           has_voyage=bool(os.environ.get("VOYAGE_API_KEY")))


@app.route("/api/sources")
def api_sources():
    return jsonify({"queries": list_db_queries(RERANKER_DB)})


@app.route("/api/delete_query", methods=["POST"])
def api_delete_query():
    p = request.get_json(silent=True) or {}
    source = (p.get("source") or "db").strip()
    query_id = p.get("query_id")
    if source == "upload":
        token = (p.get("upload_id") or "").strip()
        db_path = os.path.join(UPLOAD_DIR, token + ".db")
        if not token or not os.path.exists(db_path):
            return jsonify({"error": "فایل آپلودشده پیدا نشد."}), 400
    else:
        db_path = RERANKER_DB
        if not os.path.exists(db_path):
            return jsonify({"error": "دیتابیس پیدا نشد."}), 400
    if not query_id:
        return jsonify({"error": "شناسهٔ نامعتبر."}), 400
    try:
        conn = sqlite3.connect(db_path)
        c = conn.cursor()
        c.execute("DELETE FROM chunks WHERE url_id IN (SELECT id FROM urls WHERE query_id=?)", (query_id,))
        c.execute("DELETE FROM urls WHERE query_id=?", (query_id,))
        c.execute("DELETE FROM queries WHERE id=?", (query_id,))
        conn.commit()
        conn.close()
    except sqlite3.Error as e:
        return jsonify({"error": f"خطا در حذف: {e}"}), 500
    return jsonify({"ok": True, "queries": list_db_queries(db_path)})


@app.route("/api/upload", methods=["POST"])
def api_upload():
    f = request.files.get("db")
    if not f:
        return jsonify({"error": "فایلی ارسال نشد."}), 400
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    token = uuid.uuid4().hex[:12]
    path = os.path.join(UPLOAD_DIR, token + ".db")
    f.save(path)
    try:
        qs = list_db_queries(path)
    except Exception:
        os.remove(path)
        return jsonify({"error": "فایل دیتابیس معتبر نیست."}), 400
    return jsonify({"upload_id": token, "queries": qs})


@app.route("/api/extract", methods=["POST"])
def api_extract():
    p = request.get_json(silent=True) or {}
    provider = (p.get("provider") or "jina").strip().lower()
    if provider not in ("jina", "voyage"):
        provider = "jina"
    env_key = "VOYAGE_API_KEY" if provider == "voyage" else "JINA_API_KEY"
    api_key = (p.get("embed_key") or "").strip() or os.environ.get(env_key, "")
    source = (p.get("source") or "db").strip()
    query_id = p.get("query_id")
    try:
        keyword_num = max(5, min(int(p.get("keyword_num") or KEYWORD_NUM), 60))
    except (TypeError, ValueError):
        keyword_num = KEYWORD_NUM
    try:
        max_chars = max(2000, min(int(p.get("max_chars") or DEFAULT_MAX_CHARS), 80000))
    except (TypeError, ValueError):
        max_chars = DEFAULT_MAX_CHARS

    if source == "upload":
        token = (p.get("upload_id") or "").strip()
        db_path = os.path.join(UPLOAD_DIR, token + ".db")
        if not token or not os.path.exists(db_path):
            return jsonify({"error": "فایل آپلودشده پیدا نشد، دوباره بارگذاری کنید."}), 400
    else:
        db_path = RERANKER_DB
        if not os.path.exists(db_path):
            return jsonify({"error": "دیتابیس reranker روی سرور پیدا نشد."}), 400

    label = "Voyage" if provider == "voyage" else "Jina"
    if not api_key:
        return jsonify({"error": f"کلید {label} لازم است."}), 400

    jid = new_job()
    threading.Thread(target=run_extract,
                     args=(jid, db_path, query_id, provider, api_key, keyword_num, max_chars),
                     daemon=True).start()
    return jsonify({"job_id": jid})


@app.route("/api/job/<jid>")
def api_job(jid):
    j = job_get(jid)
    if not j:
        return jsonify({"error": "job not found"}), 404
    return jsonify(j)


@app.route("/api/export/<jid>")
def api_export(jid):
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    j = job_get(jid)
    if not j or not j.get("results"):
        return jsonify({"error": "نتیجه‌ای برای خروجی نیست."}), 404
    data = j["results"]
    items = data["items"]
    head_font = Font(bold=True, name="Arial", color="FFFFFF")
    alt = PatternFill("solid", start_color="F0F4F8")
    none = PatternFill("solid", start_color="FFFFFF")

    wb = Workbook()
    ws1 = wb.active; ws1.title = "کلیدواژه‌ها به تفکیک URL"
    h1 = ["رتبه گوگل", "URL", "عنوان صفحه", "کیورد اصلی", "رتبه کلیدواژه", "کلیدواژه استخراج‌شده"]
    for c, h in enumerate(h1, 1):
        cell = ws1.cell(1, c, h); cell.font = head_font
        cell.fill = PatternFill("solid", start_color="2E4057")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r = 2
    for it in items:
        for kr, kw in enumerate(it.get("keywords", []), 1):
            fill = alt if r % 2 == 0 else none
            for c, v in enumerate([it["rank"], it["url"], it["title"], it["original_query"], kr, kw], 1):
                cell = ws1.cell(r, c, v); cell.font = Font(name="Arial", size=10)
                cell.fill = fill; cell.alignment = Alignment(wrap_text=True, vertical="center")
            r += 1
    for col, w in zip("ABCDEF", [12, 55, 40, 25, 14, 40]):
        ws1.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Top کلیدواژه‌ها")
    h2 = ["رتبه گوگل", "URL", "عنوان صفحه", "کیورد اصلی"] + [f"کلیدواژه {i}" for i in range(1, 11)]
    for c, h in enumerate(h2, 1):
        cell = ws2.cell(1, c, h); cell.font = head_font
        cell.fill = PatternFill("solid", start_color="1A535C")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for ri, it in enumerate(items, 2):
        top10 = it.get("keywords", [])[:10]
        vals = [it["rank"], it["url"], it["title"], it["original_query"]] + top10 + [""] * (10 - len(top10))
        fill = alt if ri % 2 == 0 else none
        for c, v in enumerate(vals, 1):
            cell = ws2.cell(ri, c, v); cell.font = Font(name="Arial", size=10)
            cell.fill = fill; cell.alignment = Alignment(wrap_text=True, vertical="center")
    for col, w in zip("ABCD", [12, 50, 35, 25]):
        ws2.column_dimensions[col].width = w
    for i in range(5, 15):
        ws2.column_dimensions[get_column_letter(i)].width = 26

    ws3 = wb.create_sheet("فرکوانسی کلیدواژه‌ها")
    h3 = ["کلیدواژه", "تکرار در URLها", "درصد از کل"]
    for c, h in enumerate(h3, 1):
        cell = ws3.cell(1, c, h); cell.font = head_font
        cell.fill = PatternFill("solid", start_color="4F6D7A")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    total = data["total_urls"] or 1
    for ri, (kw, cnt) in enumerate(data["freq"], 2):
        fill = alt if ri % 2 == 0 else none
        ws3.cell(ri, 1, kw).font = Font(name="Arial", size=10)
        ws3.cell(ri, 2, cnt).font = Font(name="Arial", size=10)
        pc = ws3.cell(ri, 3, f"=B{ri}/{total}"); pc.number_format = "0.0%"; pc.font = Font(name="Arial", size=10)
        for c in range(1, 4):
            ws3.cell(ri, c).fill = fill
    for col, w in zip("ABC", [40, 18, 18]):
        ws3.column_dimensions[col].width = w

    bio = io.BytesIO(); wb.save(bio); bio.seek(0)
    return send_file(bio, as_attachment=True, download_name=f"keyword_research_{jid}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/healthz")
def healthz():
    return jsonify({"status": "ok", "pos_model": os.path.exists(POS_MODEL_PATH)})


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8002, debug=True)
